# -*- coding: utf-8 -*-
"""
ساخت فایل ayani_rates.json (نرخ اعیانی هر متر مربع به ریال) از روی اکسل منبع.

ورودی‌ها:
  data/ayani_rates_source.xlsx  ← اکسل ارسالی کارفرما (شیت اول: استان‌ها و شهرستان‌ها)
  data/county_seats.json        ← مختصات تقریبی مرکز هر شهرستان (برای یافتن همسایه)

خروجی:
  ayani_rates.json  (ریشهٔ پروژه — توسط ayani_calc.py خوانده می‌شود)

قواعد:
  - اعداد اکسل «عیناً» استفاده می‌شوند (ریال / متر مربع).
  - ردیف‌هایی که در استان اشتباه ثبت شده‌اند (MISPLACED_ROWS) حذف می‌شوند.
  - هر خانهٔ خالی/غیرعددی از نزدیک‌ترین شهرستانِ همان استان که آن خانه را
    دارد پر می‌شود و منبعش در فیلد "filled" ثبت می‌شود.

اجرا (هر بار که اکسل عوض شد):
  python scripts/build_ayani_rates.py
"""

import datetime
import json
import math
import os
import sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC_XLSX = os.path.join(ROOT, "data", "ayani_rates_source.xlsx")
SEATS_JSON = os.path.join(ROOT, "data", "county_seats.json")
OUT_JSON = os.path.join(ROOT, "ayani_rates.json")

# ستون‌های اکسل (۱-مبنا): (کلید کاربری، کلید سازه) → شماره ستون
# D/E مسکونی، F/G تجاری، H/I اداری، J/K کشاورزی، L/M صنعتی/خدماتی
COLUMNS = {
    ("residential", "concrete"): 4, ("residential", "other"): 5,
    ("commercial", "concrete"): 6, ("commercial", "other"): 7,
    ("administrative", "concrete"): 8, ("administrative", "other"): 9,
    ("agricultural", "concrete"): 10, ("agricultural", "other"): 11,
    ("industrial", "concrete"): 12, ("industrial", "other"): 13,
}


# ردیف‌هایی از اکسل که در استان اشتباه ثبت شده‌اند و همان شهرستان در استان
# درست خودش ردیف جداگانه (با نرخ مخصوص خودش) دارد → ردیف اشتباه حذف می‌شود.
MISPLACED_ROWS = {
    ("سیستان و بلوچستان", "گرگان"): "گرگان مرکز استان گلستان است (ردیف گلستان/گرگان موجود است)",
    ("سیستان و بلوچستان", "مارگون"): "مارگون در کهگیلویه و بویراحمد است (ردیف آن موجود است)",
    ("گلستان", "رشتخوار"): "رشتخوار در خراسان رضوی است (ردیف آن موجود است)",
    ("آذربایجان غربی", "نوک‌آباد"): "نوک‌آباد مرکز شهرستان تفتان در سیستان و بلوچستان است (ردیف تفتان موجود است)",
}


def _num(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)) and v > 0:
        return int(v)
    if isinstance(v, str):
        s = v.strip().replace(",", "").translate(str.maketrans("۰۱۲۳۴۵۶۷۸۹", "0123456789"))
        if s.isdigit() and int(s) > 0:
            return int(s)
    return None


def _dist(a, b):
    if a["lat"] is None or b["lat"] is None:
        return float("inf")
    # فاصلهٔ تقریبی (کافی برای انتخاب نزدیک‌ترین همسایه)
    dlat = a["lat"] - b["lat"]
    dlng = (a["lng"] - b["lng"]) * math.cos(math.radians((a["lat"] + b["lat"]) / 2))
    return math.hypot(dlat, dlng)


def build():
    ws = openpyxl.load_workbook(SRC_XLSX, data_only=True).worksheets[0]
    seats = {(s["province"], s["county"]): s for s in json.load(open(SEATS_JSON, encoding="utf-8"))}

    rows = []
    for r in range(2, ws.max_row + 1):
        province = (ws.cell(r, 2).value or "").strip()
        county = (ws.cell(r, 3).value or "").strip()
        if not province or not county:
            continue
        if (province, county) in MISPLACED_ROWS:
            print(f"   ✂ حذف ردیف جابه‌جا: {province} / {county} — {MISPLACED_ROWS[(province, county)]}")
            continue
        seat = seats.get((province, county), {})
        rates = {}
        for (use, struct), col in COLUMNS.items():
            rates.setdefault(use, {})[struct] = _num(ws.cell(r, col).value)
        rows.append({
            "province": province, "county": county,
            "lat": seat.get("lat"), "lng": seat.get("lng"),
            "rates": rates, "filled": {},
        })

    # پرکردن خانه‌های خالی از نزدیک‌ترین شهرستان هم‌استان
    for row in rows:
        for use, structs in row["rates"].items():
            for struct, val in structs.items():
                if val is not None:
                    continue
                donors = [
                    d for d in rows
                    if d is not row and d["province"] == row["province"]
                    and d["rates"][use][struct] is not None
                    and not d["filled"].get(f"{use}.{struct}")
                ]
                if not donors:
                    raise SystemExit(f"هیچ همسایه‌ای برای {row['province']}/{row['county']} {use}.{struct} نیست")
                donor = min(donors, key=lambda d: _dist(row, d))
                structs[struct] = donor["rates"][use][struct]
                row["filled"][f"{use}.{struct}"] = donor["county"]

    provinces = {}
    for row in rows:
        p = row.pop("province")
        if not row["filled"]:
            row.pop("filled")
        provinces.setdefault(p, []).append(row)

    out = {
        "version": 1,
        "unit": "ریال بر متر مربع",
        "generated_at": datetime.date.today().isoformat(),
        "source": os.path.basename(SRC_XLSX),
        "provinces": provinces,
    }
    with open(OUT_JSON, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=1)

    filled = [(p, c["county"], c["filled"]) for p, cs in provinces.items() for c in cs if c.get("filled")]
    print(f"✅ {sum(len(v) for v in provinces.values())} شهرستان در {len(provinces)} استان → {OUT_JSON}")
    for p, c, f in filled:
        print(f"   ↳ پر شده: {p} / {c}: {f}")


if __name__ == "__main__":
    sys.exit(build())
