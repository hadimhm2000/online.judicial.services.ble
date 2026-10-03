# -*- coding: utf-8 -*-
"""
build_bulk_templates.py
──────────────────────────────────────────────────────────────────────────
ساخت قالب‌های اکسل دسته‌جمعی نسخهٔ ۲ در پوشهٔ templates از روی تعاریف
bulk_excel_v2.py. هر بار که ستون‌ها یا فهرست شعب تغییر کرد، دوباره اجرا شود:

    python build_bulk_templates.py

خروجی‌ها:
    templates/sample_lavayeh_bulk.xlsx
    templates/sample_ezhharnameh_bulk.xlsx
    templates/sample_check_bulk.xlsx
    templates/sample_bulk_inquiry.xlsx

نکته‌های فنی:
  • فرمول‌های TEXTJOIN حذف شدند (در اکسل ۲۰۱۶ و قبل‌تر کار نمی‌کنند).
  • لیست‌های چندمرحله‌ای شعبه با OFFSET/MATCH روی شیت پنهان «_لیست‌ها»
    ساخته می‌شوند: هر گزینه یک کلید دارد («ROOT|استان|حوزه|...») و لیست هر
    مرحله، بلوک پشت‌سرهم همان کلید است. در Excel و WPS کار می‌کند؛ Google
    Sheets لیست وابسته را نشان نمی‌دهد (تایپ دستی همچنان پذیرفته می‌شود).
  • ستون‌های کدملی/شماره متنی (@) هستند تا صفر اول و رقم‌های آخر از بین نروند.
"""
import os
import re
from collections import OrderedDict

import openpyxl
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

import bulk_excel_v2 as v2

TEMPLATES_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "templates")

REQUIRED_FILL = PatternFill("solid", fgColor="1F4E78")
OPTIONAL_FILL = PatternFill("solid", fgColor="BDD7EE")
SAMPLE_FILL = PatternFill("solid", fgColor="EDEDED")
BAD_FILL = PatternFill("solid", fgColor="F8C9C6")
HEADER_FONT_REQ = Font(name="Tahoma", bold=True, color="FFFFFF")
HEADER_FONT_OPT = Font(name="Tahoma", bold=True, color="1F1F1F")
SAMPLE_FONT = Font(name="Tahoma", italic=True, color="7F7F7F")
TEXT_KINDS = {"nid", "pid", "digits"}


def _nat_key(value):
    return [int(x) if x.isdigit() else x for x in re.split(r"(\d+)", value)]


def _tree_blocks(paths):
    """paths: تاپل‌های سطح‌به‌سطح → OrderedDict کلید → گزینه‌های مرتب‌شده."""
    blocks = OrderedDict()
    for path in paths:
        for i in range(len(path)):
            key = "|".join(("ROOT",) + tuple(path[:i]))
            blocks.setdefault(key, set()).add(path[i])
    ordered = OrderedDict()
    for key in sorted(blocks, key=lambda k: (k.count("|"), k)):
        ordered[key] = sorted(blocks[key], key=_nat_key)
    return ordered


class _Lists:
    """شیت پنهان «_لیست‌ها»: لیست‌های ساده در ستون‌های جدا، درخت‌ها به‌صورت کلید/مقدار."""

    def __init__(self, wb):
        self.wb = wb
        self.ws = wb.create_sheet(v2.LISTS_SHEET)
        self.ws.sheet_state = "hidden"
        self.next_col = 1

    def add_simple(self, name, values):
        col = get_column_letter(self.next_col)
        for i, v in enumerate(values, start=1):
            self.ws.cell(row=i, column=self.next_col, value=v)
        self.wb.defined_names[name] = DefinedName(
            name, attr_text=f"'{v2.LISTS_SHEET}'!${col}$1:${col}${len(values)}")
        self.next_col += 1

    def add_tree(self, name, paths):
        kcol, vcol = self.next_col, self.next_col + 1
        r = 1
        for key, values in _tree_blocks(paths).items():
            for v in values:
                self.ws.cell(row=r, column=kcol, value=key)
                self.ws.cell(row=r, column=vcol, value=v)
                r += 1
        kl, vl = get_column_letter(kcol), get_column_letter(vcol)
        self.wb.defined_names[f"{name}Key"] = DefinedName(
            f"{name}Key", attr_text=f"'{v2.LISTS_SHEET}'!${kl}$1:${kl}${r - 1}")
        self.wb.defined_names[f"{name}Val"] = DefinedName(
            f"{name}Val", attr_text=f"'{v2.LISTS_SHEET}'!${vl}$1")
        self.next_col += 2


def _tree_formula(tree, prev_letters):
    key = '"ROOT"' if not prev_letters else '"ROOT|"&' + '&"|"&'.join(f"${c}2" for c in prev_letters)
    f = f"OFFSET({tree}Val,MATCH({key},{tree}Key,0)-1,0,COUNTIF({tree}Key,{key}),1)"
    assert len(f) < 255, f   # محدودیت اکسل برای فرمول Data Validation
    return f


def _add_data_sheet(wb, spec: v2.SheetSpec, first: bool):
    ws = wb.active if first else wb.create_sheet()
    ws.title = spec.name
    ws.sheet_view.rightToLeft = True
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 42
    last = v2.LAST_DATA_ROW
    tree_letters = {}

    for ci, col in enumerate(spec.columns, start=1):
        letter = get_column_letter(ci)
        cell = ws.cell(row=1, column=ci, value=col.header)
        cell.fill = REQUIRED_FILL if col.required else OPTIONAL_FILL
        cell.font = HEADER_FONT_REQ if col.required else HEADER_FONT_OPT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        note = ("ضروری. " if col.required else "اختیاری. ") + col.note
        cell.comment = Comment(note.strip(), "ربات")
        ws.column_dimensions[letter].width = col.width
        rng = f"{letter}2:{letter}{last}"

        if col.kind in TEXT_KINDS:
            for r in range(2, last + 1):
                ws.cell(row=r, column=ci).number_format = "@"
        if col.kind == "longtext":
            for r in range(2, last + 1):
                ws.cell(row=r, column=ci).alignment = Alignment(wrap_text=True, vertical="top")

        if col.kind == "choice":
            dv = DataValidation(type="list", formula1=col.list_name, allow_blank=True,
                                showErrorMessage=True, errorTitle="انتخاب از لیست",
                                error="لطفاً یکی از گزینه‌های لیست را انتخاب کنید.")
            ws.add_data_validation(dv)
            dv.add(rng)
        elif col.kind == "tree":
            prev = tree_letters.setdefault(col.tree, [])
            dv = DataValidation(type="list", formula1=_tree_formula(col.tree, prev), allow_blank=True,
                                showErrorMessage=True, errorTitle="انتخاب از لیست",
                                error="از لیست انتخاب کنید. اول ستون‌های قبلی (استان، حوزه، ...) را پر کنید.")
            ws.add_data_validation(dv)
            dv.add(rng)
            prev.append(letter)
        elif col.kind == "nid":
            ws.conditional_formatting.add(rng, FormulaRule(
                formula=[f'AND({letter}2<>"",LEN({letter}2)<>10)'], fill=BAD_FILL))
        elif col.kind == "pid":
            ws.conditional_formatting.add(rng, FormulaRule(
                formula=[f'AND({letter}2<>"",LEN({letter}2)<>10,LEN({letter}2)<>11)'], fill=BAD_FILL))

    for ci, col in enumerate(spec.columns, start=1):
        if col.key in spec.sample:
            c = ws.cell(row=2, column=ci, value=spec.sample[col.key])
            c.fill = SAMPLE_FILL
            c.font = SAMPLE_FONT
    return ws


def _add_guide(wb, lines):
    ws = wb.create_sheet(v2.GUIDE_SHEET)
    ws.sheet_view.rightToLeft = True
    ws.column_dimensions["A"].width = 110
    for i, line in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=line)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if i == 1:
            c.font = Font(name="Tahoma", bold=True, size=13)
        else:
            c.font = Font(name="Tahoma", size=11)


def _add_meta(wb, service):
    ws = wb.create_sheet(v2.META_SHEET)
    ws["A1"] = v2.TEMPLATE_VERSION
    ws["A2"] = service
    ws.sheet_state = "hidden"


_COMMON_GUIDE = [
    "ردیف خاکستری هر شیت نمونه است. رویش بنویسید یا پاکش کنید؛ اگر دست‌نخورده بماند ربات آن را نادیده می‌گیرد.",
    "ستون‌های آبی تیره ضروری‌اند و آبی روشن اختیاری. برای دیدن توضیح هر ستون، روی عنوانش نگه دارید.",
    "ستون‌های ▼ را از لیست انتخاب کنید. خانه‌ای که کدملی‌اش ۱۰ رقم نباشد صورتی می‌شود.",
    "اگر ردیفی خطا داشت، ربات همین فایل را با خانه‌های قرمز و توضیح خطا برمی‌گرداند؛ فقط همان‌ها را اصلاح کنید و دوباره بفرستید.",
    "فایل را در پیش‌نمایش بله ویرایش نکنید؛ دانلود کنید و با Excel یا WPS باز کنید (لیست‌های مرحله‌ای در Google Sheets کار نمی‌کنند).",
]


def build_service(service, guide_lines):
    wb = openpyxl.Workbook()
    for i, spec in enumerate(v2.SHEETS_BY_SERVICE[service]):
        _add_data_sheet(wb, spec, first=(i == 0))
    _add_guide(wb, guide_lines + _COMMON_GUIDE)
    lists = _Lists(wb)
    for name, values in v2.SIMPLE_LISTS.items():
        lists.add_simple(name, values)
    if service == "lavayeh":
        lists.add_simple("CaseProvinces", v2.case_provinces())
        lists.add_tree("lav", v2.lavayeh_tree_paths())
    if service == "check":
        lists.add_tree("chk", v2.check_tree_paths())
    _add_meta(wb, service)
    wb.active = 0
    return wb


def build_inquiry():
    """قالب استعلام: همان ۵ ستون A تا E که bulk_inquiry_excel.py می‌خواند."""
    from bulk_inquiry_excel import CATEGORY_OPTIONS, DATA_SHEET_NAME, SAMPLE_ROW

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = DATA_SHEET_NAME
    ws.sheet_view.rightToLeft = True
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 42
    last = v2.LAST_DATA_ROW
    cols = [
        ("کدرهگیری (۱۶ رقم)", 24, True, "کدرهگیری ۱۶ رقمی. فقط اگر استعلام با کدرهگیری می‌خواهید."),
        ("دریافت پیوست‌ها؟ ▼", 16, False, "بله یا خیر. فقط برای کدرهگیری."),
        ("نوع سند ▼ (اگر پیوست=بله)", 36, False, "اگر پیوست می‌خواهید، نوع سند را از لیست انتخاب کنید."),
        ("شماره موبایل", 18, True, "۱۱ رقم، با ۰۹ شروع می‌شود."),
        ("کدملی", 16, True, "کدملی ۱۰ رقمی."),
    ]
    for ci, (header, width, text, note) in enumerate(cols, start=1):
        letter = get_column_letter(ci)
        c = ws.cell(row=1, column=ci, value=header)
        c.fill = REQUIRED_FILL
        c.font = HEADER_FONT_REQ
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.comment = Comment(note, "ربات")
        ws.column_dimensions[letter].width = width
        if text:
            for r in range(2, last + 1):
                ws.cell(row=r, column=ci).number_format = "@"
    for ci, v in enumerate(SAMPLE_ROW, start=1):
        if v:
            c = ws.cell(row=2, column=ci, value=v)
            c.fill = SAMPLE_FILL
            c.font = SAMPLE_FONT

    lists = wb.create_sheet("لیست‌ها")
    lists.sheet_state = "hidden"
    for i, v in enumerate(CATEGORY_OPTIONS, start=1):
        lists.cell(row=i, column=1, value=v)
    dv = DataValidation(type="list", formula1='"بله,خیر"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"B2:B{last}")
    dv = DataValidation(type="list", formula1=f"'لیست‌ها'!$A$1:$A${len(CATEGORY_OPTIONS)}", allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"C2:C{last}")
    ws.conditional_formatting.add(f"A2:A{last}", FormulaRule(
        formula=['AND(A2<>"",LEN(A2)<>16)'], fill=BAD_FILL))
    ws.conditional_formatting.add(f"D2:D{last}", FormulaRule(
        formula=['AND(D2<>"",LEN(D2)<>11)'], fill=BAD_FILL))
    ws.conditional_formatting.add(f"E2:E{last}", FormulaRule(
        formula=['AND(E2<>"",LEN(E2)<>10)'], fill=BAD_FILL))

    _add_guide(wb, [
        "راهنمای استعلام دسته‌جمعی",
        "هر ردیف یک استعلام است. فقط ستونی را پر کنید که لازم دارید: کدرهگیری، موبایل یا کدملی.",
        "برای کدرهگیری با پیوست، «دریافت پیوست‌ها؟» را «بله» بگذارید و «نوع سند» را از لیست انتخاب کنید.",
        "ستون‌ها متنی هستند، پس صفر اول موبایل و کدملی حذف نمی‌شود. گذاشتن حرف T/M/N قبل از عدد دیگر لازم نیست.",
        "اگر عددها را از فایل دیگری کپی می‌کنید، با «Paste Values» (فقط مقدار) بچسبانید تا متنی بمانند.",
    ] + _COMMON_GUIDE[:1] + _COMMON_GUIDE[3:])
    wb.active = 0
    return wb


GUIDES = {
    "lavayeh": [
        "راهنمای ثبت دسته‌جمعی لوایح",
        "هر لایحه یک ردیف است. اگر شماره پرونده دارید از شیت «با شماره پرونده» و اگر ندارید از شیت «با شماره بایگانی» استفاده کنید. می‌توانید هر دو شیت را پر کنید و یک فایل بفرستید.",
        "برای شعبه به ترتیب استان، حوزه قضایی، مرجع و شعبه را از لیست‌ها انتخاب کنید. اگر در لیستی فقط «—» بود، همان را انتخاب کنید.",
    ],
    "ezhharnameh": [
        "راهنمای ثبت دسته‌جمعی اظهارنامه",
        "هر اظهارنامه یک ردیف است. شیت را بر اساس اظهارکننده انتخاب کنید: «اظهارکننده حقیقی» برای یک نفر، «اظهارکننده حقوقی (شرکت)» وقتی اظهارکننده شرکت است، و «چند اظهارکننده» برای ۲ تا ۴ نفر یا وکیل.",
        "مخاطب در هر سه شیت ۵ ستون دارد و حداقل یکی لازم است. کدملی ۱۰ رقمی یعنی شخص حقیقی و شناسه ملی ۱۱ رقمی یعنی شرکت؛ برای مخاطب شرکت نماینده لازم نیست.",
        "می‌توانید چند شیت را هم‌زمان پر کنید و یک فایل بفرستید.",
    ],
    "check": [
        "راهنمای ثبت دسته‌جمعی دعاوی چک",
        "هر دادخواست یک ردیف است. شیت را بر اساس خواهان انتخاب کنید: «خواهان حقیقی» برای یک نفر، «خواهان حقوقی (شرکت)» وقتی خواهان شرکت است، و «چند خواهان» برای ۲ تا ۴ نفر یا وکیل.",
        "خوانده در هر سه شیت ۵ ستون دارد و حداقل یکی لازم است. کدملی ۱۰ رقمی یعنی شخص حقیقی و شناسه ملی ۱۱ رقمی یعنی شرکت؛ برای خوانده شرکت نماینده لازم نیست.",
        "بعد از ارسال فایل، ربات برای هر ردیف ۳ تصویر می‌خواهد: روی چک، پشت چک و گواهی عدم پرداخت.",
        "می‌توانید چند شیت را هم‌زمان پر کنید و یک فایل بفرستید.",
        "برای صلاحیت دادگاه به ترتیب استان، حوزه قضایی و دادگاه/مجتمع را از لیست‌ها انتخاب کنید. اگر در لیستی فقط «—» بود، همان را انتخاب کنید.",
    ],
}

OUTPUTS = {
    "lavayeh": "sample_lavayeh_bulk.xlsx",
    "ezhharnameh": "sample_ezhharnameh_bulk.xlsx",
    "check": "sample_check_bulk.xlsx",
}


def main():
    os.makedirs(TEMPLATES_DIR, exist_ok=True)
    for service, filename in OUTPUTS.items():
        path = os.path.join(TEMPLATES_DIR, filename)
        build_service(service, GUIDES[service]).save(path)
        print(f"✅ {path}")
    path = os.path.join(TEMPLATES_DIR, "sample_bulk_inquiry.xlsx")
    build_inquiry().save(path)
    print(f"✅ {path}")


if __name__ == "__main__":
    main()
