"""تست قالب‌های دسته‌جمعی نسخهٔ ۲ — اجرا: python -m pytest test_bulk_excel_v2.py

نیاز به BOT_TOKEN و ADMIN_ID دارد (مثل بقیهٔ تست‌هایی که config را import می‌کنند).
"""
import os

import openpyxl
import pytest

import bulk_excel_v2 as v2
import build_bulk_templates as builder


def _nid(prefix9: str) -> str:
    """کدملی معتبر با رقم کنترل درست."""
    digits = [int(c) for c in prefix9]
    rem = sum(d * (10 - i) for i, d in enumerate(digits)) % 11
    return prefix9 + str(rem if rem < 2 else 11 - rem)


def _cid(prefix10: str) -> str:
    """شناسه ملی ۱۱ رقمی معتبر."""
    digits = [int(c) for c in prefix10]
    dec = digits[9] + 2
    rem = sum((dec + digits[i]) * w for i, w in enumerate((29, 27, 23, 19, 17, 29, 27, 23, 19, 17))) % 11
    return prefix10 + str(0 if rem == 10 else rem)


N1, N2, N3, N4 = _nid("001234567"), _nid("004987654"), _nid("123456780"), _nid("223344556")
C1, C2 = _cid("1010234567"), _cid("1410234567")


def _fill(ws, spec, row, values):
    for key, val in values.items():
        ws.cell(row=row, column=spec.col_index(key), value=val)


@pytest.fixture(scope="module")
def templates(tmp_path_factory):
    d = tmp_path_factory.mktemp("tpl")
    paths = {}
    for service in ("lavayeh", "ezhharnameh", "check"):
        p = d / f"{service}.xlsx"
        builder.build_service(service, builder.GUIDES[service]).save(p)
        paths[service] = str(p)
    p = d / "inquiry.xlsx"
    builder.build_inquiry().save(p)
    paths["inquiry"] = str(p)
    return paths


def test_templates_have_no_textjoin_and_short_formulas(templates):
    for service in ("lavayeh", "ezhharnameh", "check"):
        wb = openpyxl.load_workbook(templates[service])
        assert v2.template_service(wb) == service
        assert wb.sheetnames[0] == v2.SHEETS_BY_SERVICE[service][0].name
        for spec in v2.SHEETS_BY_SERVICE[service]:
            ws = wb[spec.name]
            for dv in ws.data_validations.dataValidation:
                assert len(dv.formula1) < 255
            for row in ws.iter_rows(min_row=1, max_row=3, values_only=True):
                assert not any("TEXTJOIN" in str(v) for v in row if v)


def test_untouched_templates_have_no_rows(templates):
    for service in ("lavayeh", "ezhharnameh", "check"):
        res = v2.parse_v2(templates[service], service)
        assert res == {"valid_items": [], "invalid_rows": [], "total_rows": 0}


def test_lavayeh_both_sheets(templates, tmp_path):
    wb = openpyxl.load_workbook(templates["lavayeh"])
    case_spec, arch_spec = v2.LAVAYEH_SHEETS
    _fill(wb[case_spec.name], case_spec, 3, {
        "case_number": "140312345678901234", "case_province": "قم", "title": "لایحه دفاعیه",
        "p1": N1, "p2": N2, "text": "متن"})
    branch = next(p for p in v2.lavayeh_tree_paths() if p[3] != v2.NONE_MARK)
    _fill(wb[arch_spec.name], arch_spec, 3, dict(zip(("br1", "br2", "br3", "br4"), branch), **{
        "archive_number": "0123456", "title": "اعلام وکالت", "p1": N1, "text": "متن"}))
    path = tmp_path / "lav.xlsx"
    wb.save(path)

    res = v2.parse_v2(str(path), "lavayeh")
    assert res["total_rows"] == 2
    assert len(res["valid_items"]) == 1
    case = res["valid_items"][0]
    assert case["method"] == "شماره پرونده" and case["providers"] == [N1, N2]
    assert case["row_index"] == "2 (پرونده)"
    bad = res["invalid_rows"][0]
    assert bad["sheet"] == arch_spec.name and bad["excel_row"] == 3
    assert any("اعلام وکالت" in e or "وکیل" in e for e in bad["errors"])
    assert arch_spec.col_index("lawyer") in bad["cols"]


def test_lavayeh_archive_branch_code(templates, tmp_path):
    wb = openpyxl.load_workbook(templates["lavayeh"])
    spec = v2.LAVAYEH_SHEETS[1]
    paths = v2.lavayeh_tree_paths()
    branch = next(p for p in paths if p[3] != v2.NONE_MARK)
    _fill(wb[spec.name], spec, 3, dict(zip(("br1", "br2", "br3", "br4"), branch), **{
        "archive_number": "0123456", "title": "لایحه دفاعیه", "p1": N1, "text": "متن"}))
    # شعبه‌ای که در استان انتخاب‌شده نیست → رد
    other_province = next(p[0] for p in paths if p[0] != branch[0])
    _fill(wb[spec.name], spec, 4, dict(zip(("br1", "br2", "br3", "br4"), (other_province,) + branch[1:]), **{
        "archive_number": "0123456", "title": "لایحه دفاعیه", "p1": N1, "text": "متن"}))
    path = tmp_path / "arch.xlsx"
    wb.save(path)

    res = v2.parse_v2(str(path), "lavayeh")
    item = res["valid_items"][0]
    full = " / ".join(p for p in branch[1:] if p != v2.NONE_MARK)
    assert item["branch_name"] == full
    assert item["branch_code"] == v2._load_lavayeh_data()["name_to_code"][full]
    assert item["province"] == branch[0] and item["row_index"] == 2
    assert res["invalid_rows"][0]["excel_row"] == 4


def test_every_lavayeh_branch_resolves():
    for path in v2.lavayeh_tree_paths():
        code, _ = v2.resolve_lavayeh_branch(path)
        assert code, path


def test_ezhharnameh_sheets(templates, tmp_path):
    wb = openpyxl.load_workbook(templates["ezhharnameh"])
    simple, legal, multi = v2.EZHHARNAMEH_SHEETS
    _fill(wb[simple.name], simple, 3, {"d1_id": N1, "a1": N2, "text": "متن"})
    # مخاطب ۳ شرکت است (بدون نماینده) و ستون ۲ خالی مانده
    _fill(wb[legal.name], legal, 3, {"d1_type": v2.PERSON_LEGAL, "d1_id": C1, "d1_rep": N3,
                                     "a1": N2, "a3": C2, "text": "متن"})
    _fill(wb[multi.name], multi, 3, {"d1_type": v2.PERSON_LAWYER, "d1_id": N1,
                                     "a1": N2, "text": "متن"})
    # هیچ مخاطبی وارد نشده
    _fill(wb[simple.name], simple, 4, {"d1_id": N1, "text": "متن"})
    path = tmp_path / "ez.xlsx"
    wb.save(path)

    res = v2.parse_v2(str(path), "ezhharnameh")
    assert len(res["valid_items"]) == 2
    s, l = res["valid_items"]
    assert s["declarants"] == [{"type": "حقیقی", "id": N1, "company_rep": ""}]
    assert s["title"] == "سایر"
    assert l["declarants"] == [{"type": "حقوقی", "id": C1, "company_rep": N3}]
    assert l["addressees"] == [{"type": "حقیقی", "id": N2}, {"type": "حقوقی", "id": C2}]
    errors = {(r["sheet"], r["excel_row"]): r["errors"] for r in res["invalid_rows"]}
    assert "وکیل" in errors[(multi.name, 3)][0]
    assert errors[(simple.name, 4)] == ["حداقل یک مخاطب لازم است (ستون «کدملی/شناسه مخاطب ۱»)"]


def test_check_sheets(templates, tmp_path):
    wb = openpyxl.load_workbook(templates["check"])
    simple, legal, multi = v2.CHECK_SHEETS
    unit = next(p for p in v2.check_tree_paths() if p[2] != v2.NONE_MARK)
    br = dict(zip(("br1", "br2", "br3"), unit))
    common = {"title": "صدور اجرائیه چک", "amount": "1,000,000", "tracking": "1402123456789012", "text": "متن"}
    _fill(wb[simple.name], simple, 3, dict(common, pl1_id=N1, df1=N2, **br))
    # خوانده‌ها: یک حقیقی و یک شرکت بدون نماینده
    _fill(wb[legal.name], legal, 3, dict(common, pl1_type=v2.PERSON_LEGAL, pl1_id=C1, pl1_rep=N3,
                                         pl1_rep_type="مدیرعامل", df1=N2, df5=C2, **br))
    _fill(wb[multi.name], multi, 3, dict(common, pl1_type=v2.PERSON_NATURAL, pl1_id=N1,
                                         pl2_type=v2.PERSON_NATURAL, pl2_id=N1,
                                         df1=N2, **br))
    # کدرهگیری که اکسل به عدد تبدیل کرده
    _fill(wb[simple.name], simple, 4, dict(common, pl1_id=N1, df1=N2, tracking=1402123456789010, **br))
    # کد خوانده نه ۱۰ رقم معتبر است نه ۱۱ رقم
    _fill(wb[simple.name], simple, 5, dict(common, pl1_id=N1, df2="123456789012", **br))
    path = tmp_path / "ck.xlsx"
    wb.save(path)

    res = v2.parse_v2(str(path), "check")
    assert len(res["valid_items"]) == 2
    s, l = res["valid_items"]
    assert s["check_amount"] == 1000000
    assert s["check_branch_code"] and s["check_branch_path"]
    assert s["check_plainiffs"][0]["national_id"] == N1
    assert s["check_khasteh_text"] == v2.CHECK_DEFAULT_KHASTEH["صدور اجرائیه چک"]
    assert l["check_plainiffs"][0] == {"person_type": "شخص حقوقی", "company_id": C1,
                                       "representative_type": "مدیرعامل", "national_id": N3}
    assert l["check_defendants"] == [
        {"person_type": "شخص حقیقی", "national_id": N2, "name": "---", "representative_type": ""},
        {"person_type": "شخص حقوقی", "company_id": C2, "representative_type": "", "national_id": ""},
    ]
    errors = {r["excel_row"]: r for r in res["invalid_rows"]}
    assert any("تکراری" in e for e in errors[3]["errors"])  # چند نفر: کدملی تکراری
    assert any("رقم آخر" in e for e in errors[4]["errors"])
    assert errors[5]["errors"] == ["کد خوانده 2 «123456789012» معتبر نیست (حقیقی: کدملی ۱۰ رقمی، شرکت: شناسه ملی ۱۱ رقمی)"]


def test_partially_edited_sample_is_read(templates, tmp_path):
    wb = openpyxl.load_workbook(templates["ezhharnameh"])
    spec = v2.EZHHARNAMEH_SHEETS[0]
    _fill(wb[spec.name], spec, 2, {"text": "متن واقعی"})
    path = tmp_path / "s.xlsx"
    wb.save(path)
    res = v2.parse_v2(str(path), "ezhharnameh")
    # کدملی‌های نمونه معتبر نیستند، پس ردیف نیمه‌ویرایش‌شده بی‌صدا ثبت نمی‌شود
    assert res["total_rows"] == 1 and not res["valid_items"]


def test_error_workbook(templates, tmp_path):
    wb = openpyxl.load_workbook(templates["ezhharnameh"])
    spec = v2.EZHHARNAMEH_SHEETS[0]
    _fill(wb[spec.name], spec, 3, {"d1_id": "123", "a1": N2, "text": "متن"})
    src = tmp_path / "src.xlsx"
    wb.save(src)
    res = v2.parse_v2(str(src), "ezhharnameh")
    out = tmp_path / "out.xlsx"
    assert v2.write_error_workbook(str(src), res["invalid_rows"], str(out))

    ws = openpyxl.load_workbook(out)[spec.name]
    err_col = len(spec.columns) + 1
    assert ws.cell(row=1, column=err_col).value == v2.ERROR_HEADER
    assert "اظهارکننده" in ws.cell(row=3, column=err_col).value
    assert ws.cell(row=3, column=spec.col_index("d1_id")).fill.fgColor.rgb.endswith("F28B82")
    # فایل برگشتی دوباره قابل خواندن است و ستون خطا نادیده گرفته می‌شود
    again = v2.parse_v2(str(out), "ezhharnameh")
    assert again["total_rows"] == 1
    # بعد از اصلاح، خطای قبلی پاک می‌شود
    wb = openpyxl.load_workbook(out)
    wb[spec.name].cell(row=3, column=spec.col_index("d1_id"), value=N1)
    wb[spec.name].cell(row=4, column=spec.col_index("text"), value="بدون اشخاص")
    wb.save(out)
    res = v2.parse_v2(str(out), "ezhharnameh")
    out2 = tmp_path / "out2.xlsx"
    assert v2.write_error_workbook(str(out), res["invalid_rows"], str(out2))
    ws = openpyxl.load_workbook(out2)[spec.name]
    assert ws.cell(row=3, column=err_col).value is None
    assert ws.cell(row=4, column=err_col).value


def test_parse_excel_file_routes_v2_and_keeps_legacy(templates, tmp_path):
    from bulk_submissions import parse_excel_file

    wb = openpyxl.load_workbook(templates["ezhharnameh"])
    spec = v2.EZHHARNAMEH_SHEETS[0]
    _fill(wb[spec.name], spec, 3, {"d1_id": N1, "a1": N2, "text": "متن"})
    path = tmp_path / "v2.xlsx"
    wb.save(path)
    assert len(parse_excel_file(str(path), "ezhharnameh")["valid_items"]) == 1

    # قالب قدیمی (legacy، بدون شیت _قالب) همچنان با پارسر قبلی خوانده می‌شود
    old = openpyxl.Workbook()
    ws = old.active
    ws.append(["ردیف", "کدملی اظهارکننده", "کدملی مخاطب", "نماینده", "عنوان", "متن"])
    ws.append([1, N1, N2, "", "", "متن"])
    ws.append([2, "", N2, "", "", "متن"])
    path = tmp_path / "old.xlsx"
    old.save(path)
    res = parse_excel_file(str(path), "ezhharnameh")
    assert len(res["valid_items"]) == 1
    assert res["invalid_rows"][0]["excel_row"] == 3 and res["invalid_rows"][0]["sheet"] == ws.title


def test_legacy_lavayeh_archive_needs_exact_branch(tmp_path):
    from bulk_submissions import parse_excel_file

    branch = next(p for p in v2.lavayeh_tree_paths() if p[3] != v2.NONE_MARK)
    full = " / ".join(p for p in branch[1:] if p != v2.NONE_MARK)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ثبت دسته‌جمعی لوایح"
    ws.append(["ردیف", "روش شماره‌گذاری ▼"] + [""] * 15)
    for name in (full, full[:25]):
        ws.append([1, "شعبه و شماره بایگانی", "", "", "", branch[0], name, "0123456",
                   "لایحه دفاعیه", N1, "", "", "", "خیر", "", "متن"])
    path = tmp_path / "old_lav.xlsx"
    wb.save(path)
    res = parse_excel_file(str(path), "lavayeh")
    assert res["valid_items"][0]["branch_code"] == v2._load_lavayeh_data()["name_to_code"][full]
    assert "پیدا نشد" in res["invalid_rows"][0]["errors"][0]


def test_inquiry_template_and_parser(templates, tmp_path):
    from bulk_inquiry_excel import build_bulk_inquiry_items

    assert build_bulk_inquiry_items(templates["inquiry"])["total_rows"] == 0  # فقط ردیف نمونه
    wb = openpyxl.load_workbook(templates["inquiry"])
    ws = wb.worksheets[0]
    ws["D3"] = 9123456789.0          # موبایل عددی بدون صفر اول
    ws["E4"] = 12345678.0            # کدملی عددی که اکسل اعشاری ذخیره کرده
    ws["A5"] = 1405220948201280      # کدرهگیری ۱۶ رقمی عددی (اکسل رقم آخر را صفر کرده)
    ws["A6"] = 1405220948201281      # عددی ولی رقم آخرش سالم مانده → پذیرفته می‌شود
    path = tmp_path / "inq.xlsx"
    wb.save(path)
    res = build_bulk_inquiry_items(str(path))
    phones = [i["phone"] for i in res["valid_items"] if i["kind"] == "phone"]
    nids = [i["national_id"] for i in res["valid_items"] if i["kind"] == "national_id"]
    assert phones == ["09123456789"]
    assert nids == ["0012345678"]
    assert "رقم آخر" in res["invalid_cells"][0]["error"]
    assert [i["tracking_code"] for i in res["valid_items"] if i["kind"] == "tracking"] == ["1405220948201281"]
