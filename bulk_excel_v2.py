# -*- coding: utf-8 -*-
"""
bulk_excel_v2.py
──────────────────────────────────────────────────────────────────────────
قالب‌های اکسل دسته‌جمعی نسخهٔ ۲ (لوایح، اظهارنامه، دعاوی چک): تعریف ستون‌ها،
خواندن فایل پرشده و ساخت «فایل خطادار» برای برگرداندن به کاربر.

چه چیزی نسبت به قالب‌های قبلی عوض شده؟
  • هر فایل چند شیت دارد و کاربر هر پرونده را در شیتی می‌نویسد که با آن جور
    است (مثلاً «۱ نفر - حقیقی»، «۱ نفر - با شخص حقوقی»، «چند نفر»). در یک
    فایل می‌شود چند شیت را هم‌زمان پر کرد؛ همهٔ شیت‌ها با هم خوانده می‌شوند.
  • ستون‌ها با «عنوان ستون» خوانده می‌شوند، نه با حرف ستون.
  • شعبه با لیست کشویی چندمرحله‌ای انتخاب می‌شود (استان ← حوزه ← ...) و
    ربات از همین انتخاب‌ها نام و کد دقیق شعبه را می‌سازد؛ اگر پیدا نشود،
    ردیف رد می‌شود (حدس زدن شعبهٔ «شبیه» ممنوع).
  • ردیف ۲ هر شیت «نمونه» است و اگر دست‌نخورده بماند نادیده گرفته می‌شود.
  • خروجی پارسرها دقیقاً همان ساختاری است که پارسرهای قبلی تولید می‌کردند،
    پس صف پردازش و سناریوها تغییری لازم ندارند.

قالب‌ها با build_bulk_templates.py از روی همین تعاریف ساخته می‌شوند.
فایل‌های پرشده با قالب‌های قدیمی همچنان با پارسرهای قبلی خوانده می‌شوند
(template_service() برای آن‌ها None برمی‌گرداند).
"""
import csv
import json
import logging
import os
import re
from collections import OrderedDict

import openpyxl
from openpyxl.styles import Font, PatternFill

from id_validation import is_valid_national_id, is_valid_legal_id

logger = logging.getLogger(__name__)

_BASE_DIR = os.path.dirname(os.path.abspath(__file__))

TEMPLATE_VERSION = "bulk-v2"
META_SHEET = "_قالب"
LISTS_SHEET = "_لیست‌ها"
GUIDE_SHEET = "راهنما"
NONE_MARK = "—"           # «این سطح ندارد» در لیست‌های چندمرحله‌ای شعبه
FIRST_DATA_ROW = 2        # ردیف ۲ = نمونه، داده از ردیف ۲ به بعد خوانده می‌شود
LAST_DATA_ROW = 1001      # لیست‌های کشویی و فرمت متنی تا این ردیف اعمال می‌شوند
ERROR_HEADER = "❌ خطا (ربات پر می‌کند)"

PERSON_NATURAL = "شخص حقیقی"
PERSON_LEGAL = "شخص حقوقی"
PERSON_LAWYER = "وکیل"

LAVAYEH_TITLES = [
    "لایحه دفاعیه", "صدور اجرائیه", "اعتراض به نظر کارشناس",
    "اعتراض به قرار رد دفتر", "اعلام وکالت",
    "درخواست ممنوعیت از خروج کشور", "درخواست کپی از مدارک پرونده",
    "درخواست مطالعه پرونده", "سایر عناوین",
]
CHECK_TITLES = ["صدور اجرائیه چک", "مطالبه وجه چک"]
REPRESENTATIVE_TYPES = ["مدیرعامل", "نماینده"]

# متن پیش‌فرض «عنوان خواسته» — همان متنی که check_handlers.py در ثبت دسته‌جمعی
# قدیمی و ثبت تکی پیشنهاد می‌دهد.
CHECK_DEFAULT_KHASTEH = {
    "صدور اجرائیه چک": (
        "به موجب یک فقره چک به شماره ... مورخ ... به عهده بانک ملی "
        "به مبلغ ... ریال با کدرهگیری ... به انضمام کلیه خسارات دادرسی و حق الوکاله وکیل "
        "و خسارات تاخيرتاديه از زمان سررسيد لغايت زمان كامل اجراي حكم و حق الوكاله وكيل"
    ),
    "مطالبه وجه چک": (
        "به موجب ........ فقره چک به شماره ......... مورخ ......... به عهده بانک ....... "
        "به انضمام کلیه هزینه های دادرسی و خسارات تاخیرتادیه از زمان سررسید "
        "لغایت زمان کامل اجرای حکم و حق الوکاله وکیل"
    ),
}


# ══════════════════════════════════════════════════════════════════════
# تعریف ستون‌ها و شیت‌ها
# ══════════════════════════════════════════════════════════════════════
class Col:
    """یک ستون قالب.

    kind:
      nid      کدملی ۱۰ رقمی (متنی)
      pid      کدملی یا شناسه ملی (بسته به نوع شخص، ۱۰ یا ۱۱ رقم)
      digits   عدد طولانی مثل شماره پرونده/کدرهگیری (متنی)
      amount   مبلغ
      choice   لیست کشویی ساده (list_name)
      tree     لیست کشویی چندمرحله‌ای شعبه (tree, level)
      text     متن کوتاه
      longtext متن بلند
    """

    def __init__(self, key, header, kind="text", required=False, width=18,
                 note="", list_name=None, tree=None, level=0):
        self.key = key
        self.header = header
        self.kind = kind
        self.required = required
        self.width = width
        self.note = note
        self.list_name = list_name
        self.tree = tree
        self.level = level


class SheetSpec:
    def __init__(self, name, short, columns, sample):
        self.name = name
        self.short = short          # برای برچسب ردیف وقتی چند شیت پر شده
        self.columns = columns
        self.sample = sample        # key -> مقدار ردیف نمونه

    def col_index(self, key):
        for i, c in enumerate(self.columns, start=1):
            if c.key == key:
                return i
        return None


# ── ستون‌های پرتکرار ─────────────────────────────────────────────────
def _nid(key, header, required=False, note=""):
    return Col(key, header, "nid", required, 16, note or "کدملی ۱۰ رقمی.")


def _person_cols(prefix, label, n, types_list, with_rep_type, required):
    """ستون‌های یک شخص در شیت‌های حقوقی/چندنفره: نوع، کد، نمایندهٔ شرکت (+ نوع نماینده)."""
    tag = "" if n is None else " " + str(n).translate(str.maketrans("0123456789", "۰۱۲۳۴۵۶۷۸۹"))
    cols = [
        Col(f"{prefix}_type", f"نوع {label}{tag} ▼", "choice", required, 14,
            "از لیست انتخاب کنید.", list_name=types_list),
        Col(f"{prefix}_id", f"کدملی/شناسه {label}{tag}", "pid", required, 17,
            "شخص حقیقی یا وکیل: کدملی ۱۰ رقمی. شخص حقوقی: شناسه ملی ۱۱ رقمی."),
        Col(f"{prefix}_rep", f"کدملی نماینده شرکت {label}{tag}", "nid", False, 17,
            "فقط اگر «شخص حقوقی» است: کدملی مدیرعامل یا نماینده شرکت."),
    ]
    if with_rep_type:
        cols.append(Col(f"{prefix}_rep_type", f"سمت نماینده {label}{tag} ▼", "choice", False, 14,
                        "فقط اگر «شخص حقوقی» است.", list_name="RepTypes"))
    return cols


_LAV_BRANCH_COLS = [
    Col("br1", "استان ▼", "tree", True, 30, "اول استان را انتخاب کنید.", tree="lav", level=1),
    Col("br2", "حوزه قضایی ▼", "tree", True, 34, "بعد از انتخاب استان، حوزه را انتخاب کنید.", tree="lav", level=2),
    Col("br3", "مرجع ▼", "tree", True, 40, "دادگاه، دادسرا یا شورا. اگر «—» تنها گزینه است، همان را انتخاب کنید.", tree="lav", level=3),
    Col("br4", "شعبه ▼", "tree", True, 55, "شعبهٔ رسیدگی‌کننده. اگر «—» تنها گزینه است، همان را انتخاب کنید.", tree="lav", level=4),
]
_CHECK_BRANCH_COLS = [
    Col("br1", "استان دادگاه ▼", "tree", True, 30, "صلاحیت دادگاه: اول استان را انتخاب کنید.", tree="chk", level=1),
    Col("br2", "حوزه قضایی ▼", "tree", True, 34, "بعد حوزهٔ قضایی را انتخاب کنید.", tree="chk", level=2),
    Col("br3", "دادگاه / مجتمع ▼", "tree", True, 55, "دادگاه یا مجتمع صالح. اگر «—» تنها گزینه است، همان را انتخاب کنید.", tree="chk", level=3),
]

_LAV_COMMON_TAIL = [
    Col("title", "عنوان لایحه ▼", "choice", True, 24, "از لیست انتخاب کنید.", list_name="LavTitles"),
    _nid("p1", "کدملی ارائه‌دهنده", True),
    _nid("lawyer", "کدملی وکیل (اختیاری)", False, "اگر لایحه را وکیل ارائه می‌دهد. برای «اعلام وکالت» الزامی است."),
    Col("text", "متن لایحه", "longtext", True, 60, "متن کامل لایحه را اینجا بچسبانید."),
    _nid("p2", "کدملی ارائه‌دهنده ۲ (اختیاری)"),
    _nid("p3", "کدملی ارائه‌دهنده ۳ (اختیاری)"),
    _nid("p4", "کدملی ارائه‌دهنده ۴ (اختیاری)"),
]

LAVAYEH_SHEETS = [
    SheetSpec("با شماره پرونده", "پرونده", [
        Col("case_number", "شماره پرونده", "digits", True, 22, "۱۸ رقم (پرونده‌های ۱۴۰۰ به بعد) یا ۱۶ رقم."),
        Col("sub_row", "ردیف فرعی (اختیاری)", "digits", False, 12, "اگر خالی بماند ۱ در نظر گرفته می‌شود."),
        Col("case_province", "استان پرونده ▼", "choice", True, 26, "از لیست انتخاب کنید.", list_name="CaseProvinces"),
    ] + _LAV_COMMON_TAIL, {
        "case_number": "140000000000000000", "sub_row": "1", "case_province": "قم",
        "title": "لایحه دفاعیه", "p1": "0000000001", "text": "متن نمونه — این ردیف را پاک کنید یا رویش بنویسید",
    }),
    SheetSpec("با شماره بایگانی", "بایگانی", _LAV_BRANCH_COLS + [
        Col("archive_number", "شماره بایگانی", "digits", True, 16, "شماره بایگانی پرونده در شعبه."),
    ] + _LAV_COMMON_TAIL, {
        "archive_number": "0000000", "title": "لایحه دفاعیه", "p1": "0000000001",
        "text": "متن نمونه — این ردیف را پاک کنید یا رویش بنویسید",
    }),
]

_EZ_TAIL = [
    _nid("case_rep1", "کدملی نماینده پرونده (اختیاری)"),
    Col("title", "عنوان اظهارنامه (اختیاری)", "text", False, 24, "اگر خالی بماند «سایر» ثبت می‌شود."),
    Col("text", "متن اظهارنامه", "longtext", True, 60, "متن کامل اظهارنامه را اینجا بچسبانید."),
]

EZHHARNAMEH_SHEETS = [
    SheetSpec("۱ نفر - حقیقی", "حقیقی", [
        _nid("d1_id", "کدملی اظهارکننده", True),
        _nid("a1_id", "کدملی مخاطب", True),
    ] + _EZ_TAIL, {
        "d1_id": "0000000001", "a1_id": "0000000002",
        "text": "متن نمونه — این ردیف را پاک کنید یا رویش بنویسید",
    }),
    SheetSpec("۱ نفر - با شخص حقوقی", "حقوقی",
              _person_cols("d1", "اظهارکننده", None, "TwoTypes", False, True)
              + _person_cols("a1", "مخاطب", None, "TwoTypes", False, True)[:2]
              + _EZ_TAIL, {
        "d1_type": PERSON_LEGAL, "d1_id": "10000000000", "d1_rep": "0000000001",
        "a1_type": PERSON_NATURAL, "a1_id": "0000000002",
        "text": "متن نمونه — این ردیف را پاک کنید یا رویش بنویسید",
    }),
    SheetSpec("چند نفر (تا ۴)", "چند نفر",
              sum((_person_cols(f"d{n}", "اظهارکننده", n, "DeclTypes", False, n == 1) for n in range(1, 5)), [])
              + sum((_person_cols(f"a{n}", "مخاطب", n, "TwoTypes", False, n == 1)[:2] for n in range(1, 5)), [])
              + [_nid("case_rep1", "کدملی نماینده پرونده ۱ (اختیاری)"),
                 _nid("case_rep2", "کدملی نماینده پرونده ۲ (اختیاری)")]
              + _EZ_TAIL[1:], {
        "d1_type": PERSON_NATURAL, "d1_id": "0000000001",
        "d2_type": PERSON_NATURAL, "d2_id": "0000000003",
        "a1_type": PERSON_NATURAL, "a1_id": "0000000002",
        "text": "متن نمونه — این ردیف را پاک کنید یا رویش بنویسید",
    }),
]

_CHECK_HEAD = [
    Col("title", "عنوان خواسته ▼", "choice", True, 20, "از لیست انتخاب کنید.", list_name="CheckTitles"),
    Col("amount", "مبلغ چک (ریال)", "amount", True, 18, "فقط عدد، به ریال."),
    Col("tracking", "کدرهگیری چک", "digits", True, 22, "کدرهگیری ثبت چک در سامانه صیاد."),
]
_CHECK_TAIL = _CHECK_BRANCH_COLS + [
    Col("text", "شرح خواسته / متن دادخواست", "longtext", True, 60, "متن کامل دادخواست."),
    Col("khasteh", "عنوان خواسته - متن کوتاه (اختیاری)", "text", False, 30, "اگر خالی بماند متن پیش‌فرض ثبت تکی گذاشته می‌شود."),
    Col("other", "سایر دلایل (اختیاری)", "text", False, 24, ""),
    _nid("w1", "کدملی گواه ۱ (اختیاری)"),
    _nid("w2", "کدملی گواه ۲ (اختیاری)"),
]
_CHECK_SAMPLE_COMMON = {
    "title": "مطالبه وجه چک", "amount": "500000000", "tracking": "0000000000000000",
    "text": "متن نمونه — این ردیف را پاک کنید یا رویش بنویسید",
}

CHECK_SHEETS = [
    SheetSpec("۱ نفر - حقیقی", "حقیقی", _CHECK_HEAD + [
        _nid("pl1_id", "کدملی خواهان", True),
        _nid("df1_id", "کدملی خوانده", True),
    ] + _CHECK_TAIL, dict(_CHECK_SAMPLE_COMMON, pl1_id="0000000001", df1_id="0000000002")),
    SheetSpec("۱ نفر - با شخص حقوقی", "حقوقی", _CHECK_HEAD
              + _person_cols("pl1", "خواهان", None, "TwoTypes", True, True)
              + _person_cols("df1", "خوانده", None, "TwoTypes", True, True)
              + _CHECK_TAIL, dict(_CHECK_SAMPLE_COMMON,
                                  pl1_type=PERSON_LEGAL, pl1_id="10000000000", pl1_rep="0000000001",
                                  pl1_rep_type="مدیرعامل", df1_type=PERSON_NATURAL, df1_id="0000000002")),
    SheetSpec("چند نفر (تا ۴)", "چند نفر", _CHECK_HEAD
              + sum((_person_cols(f"pl{n}", "خواهان", n, "DeclTypes", True, n == 1) for n in range(1, 5)), [])
              + sum((_person_cols(f"df{n}", "خوانده", n, "TwoTypes", True, n == 1) for n in range(1, 5)), [])
              + _CHECK_TAIL, dict(_CHECK_SAMPLE_COMMON,
                                  pl1_type=PERSON_NATURAL, pl1_id="0000000001",
                                  pl2_type=PERSON_NATURAL, pl2_id="0000000003",
                                  df1_type=PERSON_NATURAL, df1_id="0000000002")),
]

SHEETS_BY_SERVICE = {
    "lavayeh": LAVAYEH_SHEETS,
    "ezhharnameh": EZHHARNAMEH_SHEETS,
    "check": CHECK_SHEETS,
}

SIMPLE_LISTS = {
    "LavTitles": LAVAYEH_TITLES,
    "CheckTitles": CHECK_TITLES,
    "TwoTypes": [PERSON_NATURAL, PERSON_LEGAL],
    "DeclTypes": [PERSON_NATURAL, PERSON_LEGAL, PERSON_LAWYER],
    "RepTypes": REPRESENTATIVE_TYPES,
    # CaseProvinces از data/lavayeh_branch_provinces.json پر می‌شود
}


# ══════════════════════════════════════════════════════════════════════
# داده‌های شعبه
# ══════════════════════════════════════════════════════════════════════
_LAV_DATA = None
_CHK_ROWS = None


def _load_lavayeh_data():
    """{"provinces": {استان: [کد شعبه...]}, "case_provinces": [...], "code_to_name": {...}}"""
    global _LAV_DATA
    if _LAV_DATA is None:
        with open(os.path.join(_BASE_DIR, "data", "lavayeh_branch_provinces.json"), encoding="utf-8") as f:
            data = json.load(f)
        with open(os.path.join(_BASE_DIR, "branch_code_lookup.json"), encoding="utf-8") as f:
            lookup = json.load(f)
        data["name_to_code"] = lookup
        data["code_to_name"] = {code: name for name, code in lookup.items()}
        _LAV_DATA = data
    return _LAV_DATA


def case_provinces():
    return _load_lavayeh_data()["case_provinces"]


def lavayeh_tree_paths():
    """هر شعبه → (استان، حوزه، مرجع، شعبه). نام کامل = بخش‌های غیر «—» با « / »."""
    data = _load_lavayeh_data()
    out = []
    for province, codes in data["provinces"].items():
        for code in codes:
            parts = data["code_to_name"][code].split(" / ")
            hozeh = parts[0]
            marja = parts[1] if len(parts) > 1 else NONE_MARK
            branch = " / ".join(parts[2:]) if len(parts) > 2 else NONE_MARK
            out.append((province, hozeh, marja, branch))
    return out


def _fa(text):
    return (text or "").replace("ي", "ی").replace("ك", "ک")


def check_tree_paths():
    """هر واحد units_output.csv → (استان، حوزه، دادگاه/مجتمع)."""
    global _CHK_ROWS
    if _CHK_ROWS is None:
        rows = []
        path = os.path.join(_BASE_DIR, "units_output.csv")
        with open(path, encoding="utf-8-sig", newline="") as f:
            for row in csv.DictReader(f):
                parts = [_fa(p.strip()) for p in (row.get("Path") or "").split(" > ") if p.strip()]
                if not parts:
                    continue
                hozeh = parts[1] if len(parts) > 1 else NONE_MARK
                unit = " > ".join(parts[2:]) if len(parts) > 2 else NONE_MARK
                rows.append((parts[0], hozeh, unit))
        _CHK_ROWS = rows
    return _CHK_ROWS


def resolve_lavayeh_branch(levels):
    """(استان، حوزه، مرجع، شعبه) → (کد، نام کامل) یا ("", "")."""
    parts = [p for p in levels[1:] if p and p != NONE_MARK]
    if not parts:
        return "", ""
    name = " / ".join(parts)
    data = _load_lavayeh_data()
    code = data["name_to_code"].get(name, "")
    if code and code not in data["provinces"].get(levels[0], []):
        return "", ""   # شعبه در استان انتخاب‌شده نیست
    return code, (name if code else "")


# ══════════════════════════════════════════════════════════════════════
# خواندن سلول‌ها
# ══════════════════════════════════════════════════════════════════════
_FA_DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")


def cell_text(value) -> str:
    """مقدار سلول → متن. عدد صحیحی که اکسل به‌شکل اعشاری ذخیره کرده (مثل
    9123456789.0) بدون «.0» برمی‌گردد تا رقم اضافه‌ای ساخته نشود."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    raw = str(value).strip()
    return "".join(c for c in raw if not (0xD800 <= ord(c) <= 0xDFFF) and ord(c) != 0xFFFD)


def digits_only(text: str) -> str:
    return re.sub(r"\D", "", (text or "").translate(_FA_DIGITS))


def maybe_lost_digits(value) -> bool:
    """عدد بیش از ۱۵ رقمی که اکسل به‌شکل «عدد» ذخیره کرده: اکسل فقط ۱۵ رقم
    معنادار نگه می‌دارد و بقیه را صفر می‌کند (کدرهگیری ۱۶ رقمی → رقم آخر ۰،
    شماره پرونده ۱۸ رقمی → سه رقم آخر ۰). اگر رقم‌های بعد از پانزدهم همه صفر
    باشند، نمی‌شود فهمید عدد اصلی چه بوده، پس باید دوباره و به‌شکل متن وارد شود."""
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return False
    try:
        digits = str(int(value))
    except (OverflowError, ValueError):
        return False
    return len(digits) > 15 and set(digits[15:]) == {"0"}


def _norm_header(text) -> str:
    t = str(text or "").replace("‌", " ").replace("‏", "").replace("‎", "")
    return re.sub(r"\s+", " ", t).strip()


def template_service(wb):
    """اگر فایل از قالب‌های نسخهٔ ۲ است، نام سرویس (lavayeh/ezhharnameh/check) را برمی‌گرداند."""
    if META_SHEET not in wb.sheetnames:
        return None
    ws = wb[META_SHEET]
    if cell_text(ws["A1"].value) != TEMPLATE_VERSION:
        return None
    service = cell_text(ws["A2"].value)
    return service if service in SHEETS_BY_SERVICE else None


def _is_sample_row(values: dict, spec: SheetSpec) -> bool:
    filled = {k: v for k, v in values.items() if v}
    return bool(filled) and filled == {k: str(v) for k, v in spec.sample.items()}


def _iter_rows(wb, service):
    """برمی‌گرداند: لیست (spec، ردیف اکسل، مقادیر متنی، مقادیر خام، نگاشت key→ستون)."""
    out = []
    for spec in SHEETS_BY_SERVICE[service]:
        if spec.name not in wb.sheetnames:
            continue
        ws = wb[spec.name]
        wanted = {_norm_header(c.header): c.key for c in spec.columns}
        colmap = {}
        for ci in range(1, ws.max_column + 1):
            key = wanted.get(_norm_header(ws.cell(row=1, column=ci).value))
            if key and key not in colmap:
                colmap[key] = ci
        for r in range(FIRST_DATA_ROW, ws.max_row + 1):
            raw = {k: ws.cell(row=r, column=ci).value for k, ci in colmap.items()}
            values = {k: cell_text(v) for k, v in raw.items()}
            if not any(values.values()):
                continue
            if _is_sample_row(values, spec):
                continue
            out.append((spec, r, values, raw, colmap))
    return out


def _row_labels(rows):
    """اگر فقط یک شیت پر شده، برچسب ردیف همان عدد ردیف است (مثل قبل)؛
    وگرنه نام کوتاه شیت هم کنارش می‌آید تا ردیف‌ها قاطی نشوند."""
    multi = len({spec.name for spec, *_ in rows}) > 1
    labels = []
    for spec, r, *_ in rows:
        labels.append(f"{r - 1} ({spec.short})" if multi else r - 1)
    return labels


class _RowErrors:
    def __init__(self):
        self.messages = []
        self.keys = []

    def add(self, msg, *keys):
        self.messages.append(msg)
        self.keys.extend(k for k in keys if k)

    def __bool__(self):
        return bool(self.messages)


def _clean_nid(text: str) -> str:
    d = digits_only(text)
    if d and len(d) < 10:
        d = d.zfill(10)   # اکسل صفرهای ابتدای کدملی را حذف کرده
    return d


def _check_nid(errs, values, key, label, required=False):
    raw = values.get(key, "")
    if not raw:
        if required:
            errs.add(f"{label} وارد نشده", key)
        return ""
    nid = _clean_nid(raw)
    if not is_valid_national_id(nid):
        errs.add(f"{label} «{raw}» کدملی معتبر نیست", key)
        return ""
    return nid


def _read_person(errs, values, prefix, label, allowed_types, need_rep_type, default_type=None):
    """یک شخص (نوع، کد، نماینده شرکت، سمت نماینده). برمی‌گرداند dict یا None."""
    ptype = values.get(f"{prefix}_type", "") or default_type or ""
    raw_id = values.get(f"{prefix}_id", "")
    if not raw_id and not values.get(f"{prefix}_type"):
        return None
    if not ptype:
        errs.add(f"نوع {label} انتخاب نشده", f"{prefix}_type")
        return None
    if ptype not in allowed_types:
        errs.add(f"نوع {label} «{ptype}» نامعتبر است (باید یکی از: {'، '.join(allowed_types)})", f"{prefix}_type")
        return None
    if not raw_id:
        errs.add(f"کدملی/شناسه {label} وارد نشده", f"{prefix}_id")
        return None
    person = {"type": ptype, "id": "", "company_rep": "", "company_rep_type": ""}
    if ptype == PERSON_LEGAL:
        cid = digits_only(raw_id)
        if not is_valid_legal_id(cid):
            errs.add(f"شناسه ملی {label} «{raw_id}» معتبر نیست (۱۱ رقم)", f"{prefix}_id")
        person["id"] = cid
        person["company_rep"] = _check_nid(errs, values, f"{prefix}_rep", f"کدملی نماینده شرکت {label}", required=True)
        if need_rep_type:
            rep_type = values.get(f"{prefix}_rep_type", "")
            if rep_type not in REPRESENTATIVE_TYPES:
                errs.add(f"سمت نماینده شرکت {label} انتخاب نشده (مدیرعامل یا نماینده)", f"{prefix}_rep_type")
            person["company_rep_type"] = rep_type
    else:
        nid = _clean_nid(raw_id)
        if not is_valid_national_id(nid):
            errs.add(f"کدملی {label} «{raw_id}» معتبر نیست", f"{prefix}_id")
        person["id"] = nid
    return person


def _plain(msg: str) -> str:
    """پیام اعتبارسنج‌های ثبت تکی (با ⚠️ و * و «مجدداً وارد کنید») → متن ساده برای اکسل."""
    msg = (msg or "").replace("⚠️", "").replace("*", "").replace("مجدداً وارد کنید:", "")
    return re.sub(r"\s+", " ", msg).strip(" .:")


def _check_duplicates(errs, ids):
    seen = set()
    for pid in ids:
        if pid and pid in seen:
            errs.add(f"کد {pid} در این ردیف تکراری است")
        seen.add(pid)


# ══════════════════════════════════════════════════════════════════════
# لوایح
# ══════════════════════════════════════════════════════════════════════
def _parse_lavayeh_row(spec, values, raw, errs):
    title = values.get("title", "")
    if not title:
        errs.add("عنوان لایحه انتخاب نشده", "title")
    elif title not in LAVAYEH_TITLES:
        errs.add(f"عنوان لایحه «{title}» نامعتبر است", "title")

    providers = []
    for n in range(1, 5):
        nid = _check_nid(errs, values, f"p{n}", f"کدملی ارائه‌دهنده {n}" if n > 1 else "کدملی ارائه‌دهنده",
                         required=(n == 1))
        if nid:
            providers.append(nid)
    lawyer = _check_nid(errs, values, "lawyer", "کدملی وکیل", required=(title == "اعلام وکالت"))
    _check_duplicates(errs, providers + [lawyer])

    text = values.get("text", "")
    if not text:
        errs.add("متن لایحه خالی است", "text")

    item = {"title": title, "providers": providers, "text": text,
            "attachments": [], "status": "pending"}
    if lawyer:
        item["lawyer_id"] = lawyer

    if spec.name == LAVAYEH_SHEETS[0].name:
        from lavayeh_handlers import validate_tracking_code
        case_number = digits_only(values.get("case_number", ""))
        if maybe_lost_digits(raw.get("case_number")):
            errs.add("شماره پرونده به‌شکل عدد ذخیره شده و رقم‌های آخرش از بین رفته؛ خانه را متنی کنید و دوباره تایپ کنید", "case_number")
        elif not case_number:
            errs.add("شماره پرونده وارد نشده", "case_number")
        else:
            ok, msg = validate_tracking_code(case_number)
            if not ok:
                errs.add(f"شماره پرونده: {_plain(msg)}", "case_number")
        province = values.get("case_province", "")
        if province not in case_provinces():
            errs.add("استان پرونده از لیست انتخاب نشده", "case_province")
        item.update({"method": "شماره پرونده", "case_number": case_number,
                     "sub_row": digits_only(values.get("sub_row", "")) or "1",
                     "province": province})
    else:
        from lavayeh_handlers import validate_archive_number
        levels = [values.get(f"br{n}", "") for n in range(1, 5)]
        code, name = resolve_lavayeh_branch(levels)
        if not all(levels):
            missing = [f"br{n}" for n in range(1, 5) if not levels[n - 1]]
            errs.add("شعبه کامل انتخاب نشده (استان، حوزه، مرجع و شعبه)", *missing)
        elif not code:
            errs.add("این ترکیب استان/حوزه/مرجع/شعبه در فهرست شعب نیست؛ دوباره از لیست‌ها انتخاب کنید",
                     "br1", "br2", "br3", "br4")
        archive = digits_only(values.get("archive_number", ""))
        if not archive:
            errs.add("شماره بایگانی وارد نشده", "archive_number")
        else:
            ok, msg = validate_archive_number(archive)
            if not ok:
                errs.add(f"شماره بایگانی: {_plain(msg)}", "archive_number")
        item.update({"method": "بایگانی", "archive_number": archive,
                     "province": levels[0], "branch_name": name, "branch_code": code})
    return item


# ══════════════════════════════════════════════════════════════════════
# اظهارنامه
# ══════════════════════════════════════════════════════════════════════
_SHORT_TYPE = {PERSON_NATURAL: "حقیقی", PERSON_LEGAL: "حقوقی", PERSON_LAWYER: "وکیل"}


def _parse_ezhharnameh_row(spec, values, raw, errs):
    simple = spec.name == EZHHARNAMEH_SHEETS[0].name
    decl_types = SIMPLE_LISTS["DeclTypes"] if spec.name == EZHHARNAMEH_SHEETS[2].name else SIMPLE_LISTS["TwoTypes"]
    default = PERSON_NATURAL if simple else None

    declarants, addressees, all_ids = [], [], []
    for n in range(1, 5):
        p = _read_person(errs, values, f"d{n}", f"اظهارکننده {n}" if n > 1 else "اظهارکننده",
                         decl_types, False, default)
        if p:
            declarants.append({"type": _SHORT_TYPE[p["type"]], "id": p["id"], "company_rep": p["company_rep"]})
            all_ids += [p["id"], p["company_rep"]]
        p = _read_person(errs, values, f"a{n}", f"مخاطب {n}" if n > 1 else "مخاطب",
                         SIMPLE_LISTS["TwoTypes"], False, default)
        if p:
            addressees.append({"type": _SHORT_TYPE[p["type"]], "id": p["id"]})
            all_ids.append(p["id"])
    if not declarants and not errs.keys.count("d1_id"):
        errs.add("اظهارکننده وارد نشده", "d1_id")
    if not addressees and not errs.keys.count("a1_id"):
        errs.add("مخاطب وارد نشده", "a1_id")
    types = [d["type"] for d in declarants]
    if "وکیل" in types and not any(t != "وکیل" for t in types):
        errs.add("اظهارکننده وکیل دارد؛ حداقل یک اظهارکنندهٔ حقیقی یا حقوقی هم لازم است")

    representatives = []
    for key in ("case_rep1", "case_rep2"):
        nid = _check_nid(errs, values, key, "کدملی نماینده پرونده")
        if nid:
            representatives.append(nid)
    _check_duplicates(errs, all_ids + representatives)

    text = values.get("text", "")
    if not text:
        errs.add("متن اظهارنامه خالی است", "text")
    return {"declarants": declarants, "addressees": addressees,
            "representatives": representatives,
            "title": values.get("title", "") or "سایر", "text": text,
            "attachments": [], "status": "pending"}


# ══════════════════════════════════════════════════════════════════════
# دعاوی چک — خروجی هم‌شکل آیتم‌های check_handlers.check_bulk_file_upload_handler
# ══════════════════════════════════════════════════════════════════════
def _check_person_out(p):
    if p["type"] == PERSON_LEGAL:
        return {"person_type": PERSON_LEGAL, "company_id": p["id"],
                "representative_type": p["company_rep_type"], "national_id": p["company_rep"]}
    return {"person_type": p["type"], "national_id": p["id"], "name": "---", "representative_type": ""}


def _parse_check_row(spec, values, raw, errs):
    from check_branches_lookup import resolve_check_branch

    title = values.get("title", "")
    if title not in CHECK_TITLES:
        errs.add("عنوان خواسته از لیست انتخاب نشده", "title")
    amount = digits_only(values.get("amount", ""))
    if not amount or int(amount) <= 0:
        errs.add("مبلغ چک نامعتبر است (فقط عدد به ریال)", "amount")
    tracking = digits_only(values.get("tracking", ""))
    if maybe_lost_digits(raw.get("tracking")):
        errs.add("کدرهگیری چک به‌شکل عدد ذخیره شده و رقم آخرش از بین رفته؛ خانه را متنی کنید و دوباره تایپ کنید", "tracking")
    elif not tracking:
        errs.add("کدرهگیری چک وارد نشده", "tracking")

    simple = spec.name == CHECK_SHEETS[0].name
    pl_types = SIMPLE_LISTS["DeclTypes"] if spec.name == CHECK_SHEETS[2].name else SIMPLE_LISTS["TwoTypes"]
    default = PERSON_NATURAL if simple else None
    plaintiffs, defendants, all_ids = [], [], []
    for n in range(1, 5):
        p = _read_person(errs, values, f"pl{n}", f"خواهان {n}" if n > 1 else "خواهان", pl_types, True, default)
        if p:
            plaintiffs.append(p)
            all_ids += [p["id"], p["company_rep"]]
        p = _read_person(errs, values, f"df{n}", f"خوانده {n}" if n > 1 else "خوانده",
                         SIMPLE_LISTS["TwoTypes"], True, default)
        if p:
            defendants.append(p)
            all_ids += [p["id"], p["company_rep"]]
    if not plaintiffs and "pl1_id" not in errs.keys:
        errs.add("خواهان وارد نشده", "pl1_id")
    if not defendants and "df1_id" not in errs.keys:
        errs.add("خوانده وارد نشده", "df1_id")

    witnesses = []
    for key, label in (("w1", "کدملی گواه ۱"), ("w2", "کدملی گواه ۲")):
        nid = _check_nid(errs, values, key, label)
        if nid:
            witnesses.append(nid)
    _check_duplicates(errs, all_ids + witnesses)

    levels = [values.get(f"br{n}", "") for n in range(1, 4)]
    branch_code, branch_name, branch_path = "", "", ""
    if not all(levels):
        missing = [f"br{n}" for n in range(1, 4) if not levels[n - 1]]
        errs.add("صلاحیت دادگاه کامل انتخاب نشده (استان، حوزه و دادگاه)", *missing)
    else:
        path = " > ".join(p for p in levels if p != NONE_MARK)
        branch_code, found, _ = resolve_check_branch(path)
        if not branch_code:
            errs.add("این ترکیب استان/حوزه/دادگاه در فهرست واحدهای قضایی نیست؛ دوباره از لیست‌ها انتخاب کنید",
                     "br1", "br2", "br3")
        else:
            branch_name, branch_path = found["name"], found["path"]

    text = values.get("text", "")
    if not text:
        errs.add("شرح خواسته (متن دادخواست) خالی است", "text")

    return {
        "check_request_title": title,
        "check_amount": int(amount) if amount else 0,
        "check_khasteh_text": values.get("khasteh", "") or CHECK_DEFAULT_KHASTEH.get(title, ""),
        "check_tracking_no": tracking,
        "check_plainiffs": [_check_person_out(p) for p in plaintiffs],
        "check_defendants": [_check_person_out(p) for p in defendants],
        "check_witnesses": [{"national_id": w} for w in witnesses],
        "check_text": text,
        "check_text_html": "",
        "check_extra_text": values.get("other", ""),
        "check_images": [],
        "check_attachment_groups": [],
        "check_branch_code": branch_code,
        "check_branch_name": branch_name,
        "check_branch_path": branch_path,
        "check_docx_file_id": None,
        "check_docx_file_name": "",
    }


_ROW_PARSERS = {
    "lavayeh": _parse_lavayeh_row,
    "ezhharnameh": _parse_ezhharnameh_row,
    "check": _parse_check_row,
}


def parse_v2(filepath: str, service: str) -> dict:
    """خروجی: {"valid_items", "invalid_rows", "total_rows"} مثل parse_excel_file.

    هر invalid_row علاوه بر row_index و errors، کلیدهای sheet / excel_row /
    cols هم دارد تا فایل خطادار ساخته شود."""
    wb = openpyxl.load_workbook(filepath, data_only=True)
    rows = _iter_rows(wb, service)
    labels = _row_labels(rows)
    valid_items, invalid_rows = [], []
    for (spec, r, values, raw, colmap), label in zip(rows, labels):
        errs = _RowErrors()
        item = _ROW_PARSERS[service](spec, values, raw, errs)
        if errs:
            invalid_rows.append({
                "row_index": label, "errors": errs.messages,
                "sheet": spec.name, "excel_row": r,
                "cols": sorted({colmap[k] for k in errs.keys if k in colmap}),
            })
        else:
            item["row_index"] = label
            valid_items.append(item)
    return {"valid_items": valid_items, "invalid_rows": invalid_rows, "total_rows": len(rows)}


def detect_and_parse_v2(filepath: str, service: str):
    """اگر فایل قالب نسخهٔ ۲ همین سرویس است نتیجهٔ parse_v2، وگرنه None."""
    try:
        wb = openpyxl.load_workbook(filepath, read_only=True)
        found = template_service(wb)
        wb.close()
    except Exception as e:
        logger.warning(f"[BULK-V2] تشخیص قالب ناموفق: {e}")
        return None
    if found != service:
        return None
    try:
        return parse_v2(filepath, service)
    except Exception as e:
        logger.error(f"[BULK-V2] خطا در خواندن فایل {filepath}: {e}", exc_info=True)
        return {"valid_items": [], "invalid_rows": [], "total_rows": 0}


# ══════════════════════════════════════════════════════════════════════
# فایل خطادار
# ══════════════════════════════════════════════════════════════════════
_ROW_FILL = PatternFill("solid", fgColor="FDE2E1")
_CELL_FILL = PatternFill("solid", fgColor="F28B82")
_NO_FILL = PatternFill(fill_type=None)


def write_error_workbook(src_path: str, invalid_rows: list, out_path: str, default_sheet: str = None) -> bool:
    """همان فایل کاربر را با ردیف‌های خطادار صورتی، خانه‌های مقصر قرمز و
    ستون «❌ خطا» ذخیره می‌کند. invalid_rows: [{"sheet", "excel_row", "cols", "errors"}].
    اگر sheet/excel_row نداشت، از default_sheet و row_index+1 استفاده می‌شود."""
    try:
        wb = openpyxl.load_workbook(src_path)
    except Exception as e:
        logger.warning(f"[BULK-V2] فایل خطادار ساخته نشد (باز نشد): {e}")
        return False

    by_sheet = OrderedDict()
    for row in invalid_rows:
        sheet = row.get("sheet") or default_sheet
        if sheet not in wb.sheetnames:
            sheet = wb.sheetnames[0]
        excel_row = row.get("excel_row")
        if excel_row is None:
            try:
                excel_row = int(row["row_index"]) + 1
            except (TypeError, ValueError, KeyError):
                continue
        # چند خطا روی یک ردیف (مثلاً استعلام که خطاها را خانه‌به‌خانه می‌دهد) یکی می‌شوند
        rows = by_sheet.setdefault(sheet, OrderedDict())
        merged = rows.setdefault(excel_row, {"errors": [], "cols": []})
        merged["errors"] += list(row.get("errors") or [])
        merged["cols"] += list(row.get("cols") or [])

    marked = 0
    first_sheet = None
    for sheet, rows in by_sheet.items():
        ws = wb[sheet]
        err_col = None
        for ci in range(1, ws.max_column + 1):
            if cell_text(ws.cell(row=1, column=ci).value) == ERROR_HEADER:
                err_col = ci
                break
        if err_col is None:
            last = 1
            for ci in range(1, ws.max_column + 1):
                if cell_text(ws.cell(row=1, column=ci).value):
                    last = ci
            err_col = last + 1
            head = ws.cell(row=1, column=err_col, value=ERROR_HEADER)
            head.font = Font(bold=True, color="B3261E")
            ws.column_dimensions[head.column_letter].width = 60
        else:
            # خطاهای ارسال قبلی را پاک کن
            for r in range(2, ws.max_row + 1):
                if ws.cell(row=r, column=err_col).value:
                    ws.cell(row=r, column=err_col).value = None
                    for ci in range(1, err_col):
                        ws.cell(row=r, column=ci).fill = _NO_FILL

        for excel_row, row in rows.items():
            for ci in range(1, err_col):
                ws.cell(row=excel_row, column=ci).fill = _ROW_FILL
            for ci in row.get("cols") or []:
                ws.cell(row=excel_row, column=ci).fill = _CELL_FILL
            msg = ws.cell(row=excel_row, column=err_col, value=" | ".join(row.get("errors") or []))
            msg.font = Font(color="B3261E")
            marked += 1
        if first_sheet is None:
            first_sheet = sheet

    if first_sheet:
        # فایل روی اولین شیتی که خطا دارد باز شود
        for ws in wb.worksheets:
            ws.sheet_view.tabSelected = False
        wb.active = wb.sheetnames.index(first_sheet)
        wb[first_sheet].sheet_view.tabSelected = True

    if not marked:
        return False
    try:
        wb.save(out_path)
    except Exception as e:
        logger.warning(f"[BULK-V2] ذخیرهٔ فایل خطادار ناموفق: {e}")
        return False
    return True


async def send_error_workbook(chat_id: int, src_path: str, invalid_rows: list,
                              default_sheet: str = None, filename: str = "فایل_با_خطاها.xlsx",
                              others_continue: bool = False):
    """فایل خطادار را می‌سازد و برای کاربر می‌فرستد. خطا در این مرحله روند اصلی را متوقف نمی‌کند.
    others_continue: ردیف‌های سالم همین فایل جداگانه ادامه پیدا می‌کنند (پس کاربر نباید
    کل فایل را دوباره بفرستد، فقط ردیف‌های اصلاح‌شده را)."""
    if not invalid_rows:
        return False
    try:
        from bale_file_sender import send_document_direct
        base, _ = os.path.splitext(src_path)
        out_path = base + "_errors.xlsx"
        if not write_error_workbook(src_path, invalid_rows, out_path, default_sheet):
            return False
        return await send_document_direct(
            chat_id, out_path, filename=filename,
            caption=("📎 همین فایل شما با علامت‌گذاری خطاهاست.\n"
                     "ردیف‌های صورتی خطا دارند؛ خانه‌های قرمز را اصلاح کنید "
                     "(توضیح هر خطا در ستون آخر آمده).\n"
                     + ("⚠️ ردیف‌های سالم همین حالا ادامه پیدا می‌کنند؛ برای جلوگیری از ثبت تکراری، "
                        "در ارسال بعدی فقط ردیف‌های اصلاح‌شده را بفرستید."
                        if others_continue else
                        "سپس فایل را دوباره بفرستید.")))
    except Exception as e:
        logger.warning(f"[BULK-V2] ارسال فایل خطادار ناموفق: {e}")
        return False
